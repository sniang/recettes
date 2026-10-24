"""Fabrique le livre de cuisine à partir des recettes du site.

Lit src/data/recettes.json et les photos de public/photos, puis écrit dans livre/ :
  - « Le livre de cuisine.pdf » (sommaire cliquable, signets),
  - « Le livre de cuisine.docx » (sommaire avec liens),
  - « Recettes (une par fichier)/ » : une recette par fichier Word, rangées par catégorie.

Les recettes sans photo de l'auteur sont illustrées par une photo libre de droits (Wikimedia
Commons), toujours accompagnée de sa mention : auteur, licence et source.

Utilisation (macOS, polices Georgia du système) :
    python3 -m venv .venv && .venv/bin/pip install -r outils/requirements.txt
    .venv/bin/python outils/livre.py
"""
import json
import os
import re
import shutil
import unicodedata

from PIL import Image

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DONNEES = os.path.join(RACINE, "src", "data", "recettes.json")
PHOTOS = os.path.join(RACINE, "public", "photos")
SORTIE = os.path.join(RACINE, "livre")

TITRE = "Le livre de cuisine"
SOUS_TITRE = "Recettes préférées de J.-C. Bergogne et M. Nasse, son épouse"
AUTEURS = "J.-C. Bergogne et M. Nasse"
AVERTISSEMENT = ("Les températures et les temps de cuisson peuvent varier en fonction "
                 "du type ou de la marque des cuisinières !")

CATEGORIES = ["Bases", "Desserts et gourmandises", "Entrées, légumes et gratins",
              "Viandes", "Poissons et fruits de mer"]
AVEC_GROUPE_DANS_NOM = {"Viandes", "Poissons et fruits de mer"}
AVEC_GROUPES = {"Desserts et gourmandises", "Entrées, légumes et gratins", "Viandes", "Poissons et fruits de mer"}


def cle_tri(texte):
    texte = unicodedata.normalize("NFKD", texte or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9 ]", "", texte)


def charger():
    donnees = json.load(open(DONNEES, encoding="utf-8"))
    ordre = donnees["ordreGroupes"]

    def rang_groupe(g):
        return (ordre.index(g), "") if g in ordre else (99, cle_tri(g))

    recettes = []
    for r in donnees["recettes"]:
        r = dict(r)
        r["_photos"] = [(os.path.join(PHOTOS, p["src"]), p["w"], p["h"]) for p in r["photos"]][:3]
        r["_credits"] = ["Photo d’illustration : %s, %s, via Wikimedia Commons (%s)" % (p["credit"], p["licence"], p["source"])
                         for p in r["photos"] if p.get("illustration")]
        recettes.append(r)
    recettes.sort(key=lambda r: (CATEGORIES.index(r["categorie"]), rang_groupe(r.get("groupe")), cle_tri(r["titre"])))
    return recettes


def sous_titre(r):
    morceaux = []
    if r.get("groupe") and r["categorie"] in AVEC_GROUPES:
        morceaux.append(r["groupe"])
    if r.get("personnes"):
        morceaux.append(r["personnes"])
    if r.get("etoiles"):
        morceaux.append("★" * r["etoiles"])
    return " · ".join(morceaux)


def nom_fichier(r):
    t = r["titre"]
    if r["categorie"] in AVEC_GROUPE_DANS_NOM and r.get("groupe") and not cle_tri(t).startswith(cle_tri(r["groupe"])):
        t = "%s - %s" % (r["groupe"], t)
    return re.sub(r'[/:\\?*"<>|]', "-", t).strip(" .")


# ----------------------------------------------------------------------------- PDF
def faire_pdf(recettes, chemin, dossier_temp):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.pdfmetrics import registerFontFamily
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (BaseDocTemplate, Flowable, Frame, Image as RLImage, KeepTogether,
                                    NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle)
    from reportlab.platypus.tableofcontents import TableOfContents

    polices = "/System/Library/Fonts/Supplemental/"
    pdfmetrics.registerFont(TTFont("G", polices + "Georgia.ttf"))
    pdfmetrics.registerFont(TTFont("G-B", polices + "Georgia Bold.ttf"))
    pdfmetrics.registerFont(TTFont("G-I", polices + "Georgia Italic.ttf"))
    pdfmetrics.registerFont(TTFont("G-BI", polices + "Georgia Bold Italic.ttf"))
    registerFontFamily("G", normal="G", bold="G-B", italic="G-I", boldItalic="G-BI")
    pdfmetrics.registerFont(TTFont("Sym", "/Library/Fonts/Arial Unicode.ttf"))

    ACCENT, TEINTE, DOUX = colors.HexColor("#8c3b1e"), colors.HexColor("#f6efe6"), colors.HexColor("#6b6158")
    st = {
        "titre": ParagraphStyle("t", fontName="G-B", fontSize=21, leading=25, textColor=ACCENT, spaceAfter=2),
        "sous": ParagraphStyle("s", fontName="G-I", fontSize=10.5, leading=14, textColor=DOUX, spaceAfter=8),
        "h": ParagraphStyle("h", fontName="G-B", fontSize=12, leading=15, textColor=ACCENT, spaceBefore=8, spaceAfter=4),
        "partie": ParagraphStyle("p", fontName="G-B", fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=2),
        "ing": ParagraphStyle("i", fontName="G", fontSize=10, leading=13, leftIndent=9, firstLineIndent=-9),
        "etape": ParagraphStyle("e", fontName="G", fontSize=10.5, leading=14.5, leftIndent=18, firstLineIndent=-18, spaceAfter=3.5),
        "intro": ParagraphStyle("in", fontName="G-I", fontSize=10.5, leading=14.5, spaceAfter=6),
        "credit": ParagraphStyle("cr", fontName="G-I", fontSize=8, leading=10, textColor=DOUX, spaceAfter=6),
        "note": ParagraphStyle("n", fontName="G-I", fontSize=10, leading=13.5, textColor=colors.HexColor("#3d3530"),
                               leftIndent=9, firstLineIndent=-9, spaceAfter=2),
        "cat": ParagraphStyle("c", fontName="G-B", fontSize=34, leading=40, textColor=ACCENT, alignment=TA_CENTER),
        "couv": ParagraphStyle("cv", fontName="G-B", fontSize=40, leading=46, textColor=ACCENT, alignment=TA_CENTER),
        "couvsous": ParagraphStyle("cs", fontName="G-I", fontSize=15, leading=20, textColor=DOUX, alignment=TA_CENTER),
        "centre": ParagraphStyle("ce", fontName="G-I", fontSize=11, leading=15, alignment=TA_CENTER, textColor=DOUX),
        "sommaire": ParagraphStyle("so", fontName="G-B", fontSize=24, leading=30, textColor=ACCENT, spaceAfter=10),
    }
    sommaire = TableOfContents(dotsMinLevel=1)
    sommaire.levelStyles = [
        ParagraphStyle("L0", fontName="G-B", fontSize=13, leading=17, textColor=ACCENT, spaceBefore=10, spaceAfter=2),
        ParagraphStyle("L1", fontName="G", fontSize=10, leading=12.5, leftIndent=10),
        ParagraphStyle("L2", fontName="G-B", fontSize=10, leading=13, leftIndent=10, spaceBefore=4, textColor=DOUX),
        ParagraphStyle("L3", fontName="G", fontSize=10, leading=12.5, leftIndent=20),
    ]

    class Repere(Flowable):
        """Repère invisible : signet du PDF + entrée du sommaire."""
        def __init__(self, cle, texte, niveau, niveau_sommaire):
            super().__init__()
            self.cle, self.texte, self.niveau, self.niveau_sommaire = cle, texte, niveau, niveau_sommaire

        def wrap(self, *a):
            return (0, 0)

        def draw(self):
            c = self.canv
            c.bookmarkPage(self.cle)
            c.addOutlineEntry(self.texte, self.cle, level=self.niveau, closed=self.niveau == 0)
            c._doctemplate.notify("TOCEntry", (self.niveau_sommaire, escape(self.texte), c.getPageNumber(), self.cle))

    class Document(BaseDocTemplate):
        def __init__(self, nom):
            super().__init__(nom, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm,
                             bottomMargin=1.8 * cm, title=TITRE, author=AUTEURS)
            cadre = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="f")
            self.addPageTemplates([PageTemplate("nue", [cadre]), PageTemplate("normale", [cadre], onPage=self.pied)])

        def pied(self, c, d):
            c.saveState()
            c.setFont("G-I", 8.5)
            c.setFillColor(DOUX)
            c.drawCentredString(A4[0] / 2, 1.0 * cm, "%s  ·  %d" % (TITRE, c.getPageNumber()))
            c.restoreState()

    def p(texte, style):
        t = escape(texte).replace("\n", "<br/>")
        t = re.sub(r"(★+)", r'<font name="Sym" color="#c08a1e">\1</font>', t)
        return Paragraph(t, st[style])

    def rangee_photos(photos, largeur):
        if not photos:
            return []
        if len(photos) == 1:
            f, w, h = photos[0]
            k = min(largeur / w, 8.5 * cm / h)
            return [RLImage(f, w * k, h * k), Spacer(1, 8)]
        ecart, haut = 0.3 * cm, 5.2 * cm
        total = sum(w * haut / h for _, w, h in photos) + ecart * (len(photos) - 1)
        if total > largeur:
            haut *= (largeur - ecart * (len(photos) - 1)) / (total - ecart * (len(photos) - 1))
        t = Table([[RLImage(f, w * haut / h, haut) for f, w, h in photos]],
                  colWidths=[w * haut / h + (ecart if i < len(photos) - 1 else 0) for i, (_, w, h) in enumerate(photos)])
        t.setStyle(TableStyle([(k, (0, 0), (-1, -1), 0) for k in ("LEFTPADDING", "RIGHTPADDING", "TOPPADDING", "BOTTOMPADDING")]))
        t.hAlign = "LEFT"
        return [t, Spacer(1, 8)]

    def encadre_ingredients(ingredients, largeur):
        cases = [p("•  " + i, "ing") for i in ingredients]
        if len(ingredients) > 5:
            moitie = (len(cases) + 1) // 2
            lignes = [[cases[k], cases[moitie + k] if moitie + k < len(cases) else ""] for k in range(moitie)]
            t = Table(lignes, colWidths=[(largeur - 0.6 * cm) / 2] * 2)
        else:
            t = Table([[c] for c in cases], colWidths=[largeur - 0.6 * cm])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), TEINTE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                               ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
                               ("TOPPADDING", (0, 0), (-1, 0), 7), ("BOTTOMPADDING", (0, -1), (-1, -1), 7)]))
        t.hAlign = "LEFT"
        return t

    doc = Document(chemin)
    L = doc.width
    flux = [NextPageTemplate("nue"), Spacer(1, 2.2 * cm), p(TITRE, "couv"), Spacer(1, 0.5 * cm),
            p(SOUS_TITRE, "couvsous"), Spacer(1, 1.2 * cm)]
    # Couverture : mosaïque de 12 photos prises à intervalles réguliers dans le livre.
    premieres = [r["_photos"][0] for r in recettes if r["_photos"]]
    mosaique = premieres[::max(1, len(premieres) // 12)][:12]
    if mosaique:
        colonne = (L - 0.6 * cm) / 4
        cases = []
        for k, (f, w, h) in enumerate(mosaique):
            cote = min(w, h * 4 / 3)
            boite = ((w - cote) / 2, (h - cote * 3 / 4) / 2, (w + cote) / 2, (h + cote * 3 / 4) / 2)
            dst = os.path.join(dossier_temp, "couverture-%d.jpg" % k)
            Image.open(f).convert("RGB").crop(tuple(int(b) for b in boite)).resize((480, 360)).save(dst, quality=75)
            cases.append(RLImage(dst, colonne - 0.2 * cm, (colonne - 0.2 * cm) * 0.75))
        t = Table([cases[i:i + 4] for i in range(0, len(cases), 4)], colWidths=[colonne] * 4)
        t.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("TOPPADDING", (0, 0), (-1, -1), 3),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        flux.append(t)
    flux += [Spacer(1, 1.2 * cm), p(AVERTISSEMENT, "centre"),
             NextPageTemplate("normale"), PageBreak(), p("Sommaire", "sommaire"), sommaire]

    for ci, cat in enumerate(CATEGORIES):
        rs = [r for r in recettes if r["categorie"] == cat]
        if not rs:
            continue
        flux += [NextPageTemplate("nue"), PageBreak(), Spacer(1, 9 * cm), Repere("cat%d" % ci, cat, 0, 0), p(cat, "cat"),
                 Spacer(1, 0.4 * cm), p("%d recettes" % len(rs), "centre"), NextPageTemplate("normale")]
        groupe_courant = None
        for r in rs:
            flux.append(PageBreak())
            niveau = 1
            if cat in AVEC_GROUPES:
                if r.get("groupe") != groupe_courant:
                    groupe_courant = r.get("groupe")
                    flux.append(Repere("g-" + r["slug"], groupe_courant or "Divers", 1, 2))
                niveau = 2
            flux.append(Repere("r-" + r["slug"], r["titre"], niveau, 3 if cat in AVEC_GROUPES else 1))
            entete = [p(r["titre"], "titre")]
            st_texte = sous_titre(r)
            entete.append(p(st_texte, "sous") if st_texte else Spacer(1, 8))
            flux.append(KeepTogether(entete + rangee_photos(r["_photos"], L) + [p(c, "credit") for c in r["_credits"]]))
            if r.get("intro"):
                flux.append(p(r["intro"], "intro"))
            for partie in r["parties"]:
                en_attente = [p(partie["nom"], "partie")] if partie.get("nom") else []
                if partie["ingredients"]:
                    flux.append(KeepTogether(en_attente + [p("Ingrédients", "h"), encadre_ingredients(partie["ingredients"], L)]))
                    en_attente = []
                if partie["etapes"]:
                    etapes = [p("%d.  %s" % (i + 1, e), "etape") for i, e in enumerate(partie["etapes"])]
                    flux.append(KeepTogether(en_attente + [p("Préparation", "h")] + etapes[:2]))
                    flux += etapes[2:]
                    en_attente = []
                flux += en_attente
            if r["notes"]:
                flux.append(KeepTogether([p("Notes", "h")] + [p("–  " + n, "note") for n in r["notes"]]))
    doc.multiBuild(flux)


# ----------------------------------------------------------------------------- Word
def outils_word():
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    ACCENT, DOUX = RGBColor(0x8C, 0x3B, 0x1E), RGBColor(0x6B, 0x61, 0x58)

    def nouveau():
        d = docx.Document()
        s = d.sections[0]
        s.page_height, s.page_width = Cm(29.7), Cm(21)
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
        normal = d.styles["Normal"]
        normal.font.name, normal.font.size = "Georgia", Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
        normal.paragraph_format.space_after = Pt(3)
        for nom, taille, couleur in (("Heading 1", 24, ACCENT), ("Heading 2", 20, ACCENT), ("Heading 3", 12.5, ACCENT),
                                     ("Heading 4", 12.5, None)):
            style = d.styles[nom]
            style.font.name, style.font.size, style.font.bold = "Georgia", Pt(taille), True
            style.font.color.rgb = couleur or RGBColor(0x2B, 0x25, 0x21)
            rpr = style.element.get_or_add_rPr()
            polices = rpr.find(qn("w:rFonts"))
            if polices is None:
                polices = OxmlElement("w:rFonts")
                rpr.append(polices)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                polices.set(qn(a), "Georgia")
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if polices.get(qn(a)) is not None:
                    del polices.attrib[qn(a)]
        d.styles["Heading 3"].paragraph_format.space_before = Pt(10)
        d.styles["Heading 3"].paragraph_format.space_after = Pt(3)
        return d

    def signet(par, nom, numero):
        debut = OxmlElement("w:bookmarkStart")
        debut.set(qn("w:id"), str(numero))
        debut.set(qn("w:name"), nom)
        fin = OxmlElement("w:bookmarkEnd")
        fin.set(qn("w:id"), str(numero))
        if par._p.pPr is not None:
            par._p.pPr.addnext(debut)
        else:
            par._p.insert(0, debut)
        par._p.append(fin)

    def lien(par, texte, ancre, gras=False, couleur=None, taille=None):
        h = OxmlElement("w:hyperlink")
        h.set(qn("w:anchor"), ancre)
        run, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
        if gras:
            rpr.append(OxmlElement("w:b"))
        c = OxmlElement("w:color")
        c.set(qn("w:val"), couleur or "2B2521")
        rpr.append(c)
        if taille:
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), str(int(taille * 2)))
            rpr.append(sz)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "none")
        rpr.append(u)
        run.append(rpr)
        t = OxmlElement("w:t")
        t.text = texte
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        h.append(run)
        par._p.append(h)

    def saut_de_page(d):
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def recette(d, r, numero_signet=None, retour=False):
        titre = d.add_heading(r["titre"], level=2)
        if numero_signet is not None:
            signet(titre, "r-" + r["slug"], numero_signet)
        st_texte = sous_titre(r)
        if st_texte:
            run = d.add_paragraph().add_run(st_texte)
            run.italic, run.font.color.rgb = True, DOUX
        if r["_photos"]:
            par = d.add_paragraph()
            if len(r["_photos"]) == 1:
                f, w, h = r["_photos"][0]
                par.add_run().add_picture(f, width=Cm(w * min(17.0 / w, 8.5 / h)))
            else:
                haut = 5.0
                total = sum(w * haut / h for _, w, h in r["_photos"]) + 0.3 * (len(r["_photos"]) - 1)
                if total > 16.8:
                    haut *= 16.8 / total
                for k, (f, w, h) in enumerate(r["_photos"]):
                    par.add_run().add_picture(f, height=Cm(haut))
                    if k < len(r["_photos"]) - 1:
                        par.add_run("  ")
        for c in r["_credits"]:
            run = d.add_paragraph().add_run(c)
            run.italic, run.font.size, run.font.color.rgb = True, Pt(8), DOUX
        if r.get("intro"):
            d.add_paragraph().add_run(r["intro"]).italic = True
        for partie in r["parties"]:
            if partie.get("nom"):
                d.add_heading(partie["nom"], level=4)
            if partie["ingredients"]:
                d.add_heading("Ingrédients", level=3)
                for i in partie["ingredients"]:
                    par = d.add_paragraph(style="List Bullet")
                    par.add_run(i)
                    par.paragraph_format.space_after = Pt(0)
            if partie["etapes"]:
                d.add_heading("Préparation", level=3)
                for k, e in enumerate(partie["etapes"]):
                    par = d.add_paragraph()
                    par.paragraph_format.left_indent, par.paragraph_format.first_line_indent = Cm(0.7), Cm(-0.7)
                    num = par.add_run("%d.  " % (k + 1))
                    num.bold, num.font.color.rgb = True, ACCENT
                    par.add_run(e)
        if r["notes"]:
            d.add_heading("Notes", level=3)
            for n in r["notes"]:
                d.add_paragraph().add_run("– " + n).italic = True
        if retour:
            par = d.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            lien(par, "↑ Retour au sommaire", "sommaire", couleur="6B6158", taille=9)

    return dict(nouveau=nouveau, signet=signet, lien=lien, saut_de_page=saut_de_page, recette=recette,
                Pt=Pt, Cm=Cm, ACCENT=ACCENT, DOUX=DOUX, CENTRE=WD_ALIGN_PARAGRAPH.CENTER)


def faire_word(recettes, chemin):
    W = outils_word()
    Pt, Cm = W["Pt"], W["Cm"]
    d = W["nouveau"]()
    for _ in range(6):
        d.add_paragraph()
    par = d.add_paragraph()
    par.alignment = W["CENTRE"]
    run = par.add_run(TITRE)
    run.bold, run.font.size, run.font.color.rgb = True, Pt(40), W["ACCENT"]
    par = d.add_paragraph()
    par.alignment = W["CENTRE"]
    run = par.add_run(SOUS_TITRE)
    run.italic, run.font.size, run.font.color.rgb = True, Pt(15), W["DOUX"]
    for _ in range(3):
        d.add_paragraph()
    par = d.add_paragraph()
    par.alignment = W["CENTRE"]
    run = par.add_run(AVERTISSEMENT)
    run.italic, run.font.color.rgb = True, W["DOUX"]
    par = d.add_paragraph()
    par.alignment = W["CENTRE"]
    run = par.add_run("Astuce : cliquez sur le nom d'une recette dans le sommaire pour y aller directement "
                      "(Cmd ou Ctrl + clic). Le volet de navigation de Word (Affichage › Volet de navigation) "
                      "affiche aussi toutes les recettes.")
    run.font.size, run.font.color.rgb = Pt(9.5), W["DOUX"]
    W["saut_de_page"](d)
    W["signet"](d.add_heading("Sommaire", level=1), "sommaire", 1)
    for ci, cat in enumerate(CATEGORIES):
        rs = [r for r in recettes if r["categorie"] == cat]
        if not rs:
            continue
        par = d.add_paragraph()
        par.paragraph_format.space_before = Pt(10)
        W["lien"](par, cat, "cat%d" % ci, gras=True, couleur="8C3B1E", taille=13)
        groupe_courant = None
        for r in rs:
            if cat in AVEC_GROUPES and r.get("groupe") != groupe_courant:
                groupe_courant = r.get("groupe")
                par = d.add_paragraph()
                par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(4), Pt(0)
                run = par.add_run(groupe_courant or "Divers")
                run.bold, run.font.size, run.font.color.rgb = True, Pt(10), W["DOUX"]
            par = d.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9 if cat in AVEC_GROUPES else 0.5)
            par.paragraph_format.space_after = Pt(0)
            W["lien"](par, r["titre"], "r-" + r["slug"], taille=10.5)
    numero = 10
    for ci, cat in enumerate(CATEGORIES):
        rs = [r for r in recettes if r["categorie"] == cat]
        if not rs:
            continue
        W["saut_de_page"](d)
        for _ in range(8):
            d.add_paragraph()
        titre = d.add_heading(cat, level=1)
        titre.alignment = W["CENTRE"]
        W["signet"](titre, "cat%d" % ci, 2 + ci)
        for r in rs:
            W["saut_de_page"](d)
            numero += 1
            W["recette"](d, r, numero_signet=numero, retour=True)
    d.core_properties.title, d.core_properties.author = TITRE, AUTEURS
    d.save(chemin)


def faire_fichiers(recettes, dossier):
    W = outils_word()
    if os.path.exists(dossier):
        shutil.rmtree(dossier)
    pris = set()
    for r in recettes:
        sous = os.path.join(dossier, "%d - %s" % (CATEGORIES.index(r["categorie"]) + 1, r["categorie"]))
        os.makedirs(sous, exist_ok=True)
        nom, k = nom_fichier(r), 2
        while (sous, nom) in pris:
            nom, k = "%s (%d)" % (nom_fichier(r), k), k + 1
        pris.add((sous, nom))
        d = W["nouveau"]()
        W["recette"](d, r)
        d.core_properties.title, d.core_properties.author = r["titre"], AUTEURS
        d.save(os.path.join(sous, nom + ".docx"))


def main():
    import tempfile
    recettes = charger()
    os.makedirs(SORTIE, exist_ok=True)
    with tempfile.TemporaryDirectory() as temp:
        faire_pdf(recettes, os.path.join(SORTIE, TITRE + ".pdf"), temp)
    faire_word(recettes, os.path.join(SORTIE, TITRE + ".docx"))
    faire_fichiers(recettes, os.path.join(SORTIE, "Recettes (une par fichier)"))
    print(len(recettes), "recettes")


if __name__ == "__main__":
    main()
